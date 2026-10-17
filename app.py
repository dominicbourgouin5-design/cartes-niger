"""
Appli web - Cartes d'Immatriculation Ambassade du Niger a Cotonou
Flask + generation Word (.docx) et PDF en ligne
"""

import os
from copy import deepcopy
from datetime import datetime, date
from io import BytesIO

from flask import Flask, request, send_file, jsonify, render_template
from PIL import Image as PILImage
import openpyxl
from docx import Document
from docx.shared import Cm, Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from lxml import etree
from reportlab.lib.colors import HexColor
from reportlab.lib.units import cm as CM
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas as pdf_canvas

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 16 * 1024 * 1024  # 16 MB max upload

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

# Image de fond incluse dans l'app
BG_IMAGE = os.path.join(BASE_DIR, "static", "carte_fond.png")

# Police Kingthings Trypewriter 2
FONT_FILE = os.path.join(BASE_DIR, "static", "KingthingsTrypewriter2.ttf")

# ── Configuration des cartes ──────────────────────────────────────────────────

FONT_NAME     = "Kingthings Trypewriter 2"   # nom de la police dans le Word
PDF_FONT_NAME = "KingthingsTrypewriter2"     # nom enregistre pour le PDF
FONT_SIZE_PT  = 14

IMG_W_PX = 1138
IMG_H_PX = 744

CARD_W_NOM_CM = 11.7
CARD_H_NOM_CM = 7.34

PAGE_W_CM = 21.0
PAGE_H_CM = 29.7

# Une imprimante ne peut pas imprimer jusqu'au bord de la feuille.
# Les cartes sont reduites si besoin pour que la colonne de 4 cartes
# tienne dans la zone imprimable (marge minimale SAFE_MARGIN_CM).
SAFE_MARGIN_CM = 0.8
# Ecart identique entre toutes les cartes (entre elles, et entre les deux colonnes)
CARD_GAP_CM = 0.3

CARDS_PER_PAGE = 6
LEFT_COUNT  = 4
RIGHT_COUNT = 2

# Taille des cartes : la colonne de 4 cartes + 3 ecarts doit tenir dans la hauteur
# en laissant au moins SAFE_MARGIN_CM en haut et en bas.
CARD_SCALE = min(1.0, (PAGE_H_CM - 2 * SAFE_MARGIN_CM - (LEFT_COUNT - 1) * CARD_GAP_CM)
                 / (LEFT_COUNT * CARD_H_NOM_CM))

CARD_W_CM = CARD_W_NOM_CM * CARD_SCALE
CARD_H_CM = CARD_H_NOM_CM * CARD_SCALE
FONT_SIZE_EFF_PT  = FONT_SIZE_PT * CARD_SCALE
FONT_SIZE_HALF_PT = round(FONT_SIZE_EFF_PT * 2)

# Marges egales haut/bas et gauche/droite : la grille est centree sur la page
TOP_MARGIN_CM = (PAGE_H_CM - (LEFT_COUNT * CARD_H_CM + (LEFT_COUNT - 1) * CARD_GAP_CM)) / 2
SIDE_MARGIN_CM = (PAGE_W_CM - (CARD_W_CM + CARD_GAP_CM + CARD_H_CM)) / 2

# Colonne gauche : 4 cartes horizontales
LEFT_X    = SIDE_MARGIN_CM
LEFT_TOPS = [TOP_MARGIN_CM + i * (CARD_H_CM + CARD_GAP_CM) for i in range(LEFT_COUNT)]

# Colonne droite : 2 cartes tournees, alignees sur le haut de la colonne gauche
RIGHT_X    = LEFT_X + CARD_W_CM + CARD_GAP_CM
RIGHT_TOPS = [TOP_MARGIN_CM + i * (CARD_W_CM + CARD_GAP_CM) for i in range(RIGHT_COUNT)]

TEXT_COLOR = "0D0D0D"

# Positions (x, y) en px de l'image de fond 1138x744.
# y est decale de +23 px par rapport aux premieres mesures : les valeurs
# etaient environ une demi-ligne au-dessus des libelles imprimes.
FIELDS = [
    ("N_ID",         8,   12),
    ("Nom",        150,  266),
    ("Prenoms",    195,  306),
    ("Fil_de",     225,  348),
    ("Et_de",      170,  390),
    ("Date",       380,  432),
    ("Lieu",       370,  472),
    ("Profession", 255,  514),
    ("Taille",     180,  554),
    ("Adresse",    215,  596),
]

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
V_NS = "urn:schemas-microsoft-com:vml"
O_NS = "urn:schemas-microsoft-com:office:office"

EMU_PER_CM = 360000

# Ratio de la hauteur du texte au-dessus de la ligne de base (police Kingthings)
PDF_TEXT_ASCENT = 0.85

pdfmetrics.registerFont(TTFont(PDF_FONT_NAME, FONT_FILE))


# ── Fonctions utilitaires ────────────────────────────────────────────────────

def cm_emu(cm):
    return int(cm * EMU_PER_CM)

def px_to_cm_x(px):
    return px * CARD_W_CM / IMG_W_PX

def px_to_cm_y(px):
    return px * CARD_H_CM / IMG_H_PX


def _format_date(val):
    """Formate une date Excel en JJ/MM/AAAA (gere datetime et texte)."""
    if isinstance(val, (datetime, date)):
        return val.strftime("%d/%m/%Y")
    s = str(val).strip()
    # Si c'est du type "1979-10-01 00:00:00", on coupe le temps
    if " 00:00:00" in s:
        s = s.replace(" 00:00:00", "")
    # Si c'est ISO "1979-10-01", convertir en JJ/MM/AAAA
    if len(s) == 10 and s[4] == '-' and s[7] == '-':
        try:
            d = datetime.strptime(s, "%Y-%m-%d")
            return d.strftime("%d/%m/%Y")
        except ValueError:
            pass
    return s


def lire_excel(excel_file):
    """Lit le fichier Excel et retourne la liste des personnes."""
    wb = openpyxl.load_workbook(excel_file)
    if "Immatriculations" in wb.sheetnames:
        ws = wb["Immatriculations"]
    else:
        ws = wb.active

    persons = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0]:
            nid_raw = str(r[0]).strip() if r[0] else ""
            nid_raw = nid_raw.replace("N° ", "N°")
            if nid_raw and not nid_raw.endswith("."):
                nid_raw += "."

            adr_raw = str(r[9]).strip().upper() if len(r) > 9 and r[9] else ""
            if adr_raw.startswith("QUARTIER "):
                adr_raw = adr_raw[len("QUARTIER "):]
            for suffix in (", BÉNIN", ", BENIN"):
                if adr_raw.endswith(suffix):
                    adr_raw = adr_raw[:-len(suffix)]
                    break

            persons.append({
                "N_ID":       nid_raw,
                "Nom":        str(r[1]).strip().upper() if len(r) > 1 and r[1] else "",
                "Prenoms":    str(r[2]).strip().upper() if len(r) > 2 and r[2] else "",
                "Fil_de":     str(r[3]).strip().upper() if len(r) > 3 and r[3] else "",
                "Et_de":      str(r[4]).strip().upper() if len(r) > 4 and r[4] else "",
                "Date":       _format_date(r[5]) if len(r) > 5 and r[5] else "",
                "Lieu":       str(r[6]).strip().upper() if len(r) > 6 and r[6] else "",
                "Profession": str(r[7]).strip().upper() if len(r) > 7 and r[7] else "",
                "Taille":     str(r[8]).strip() if len(r) > 8 and r[8] else "",
                "Adresse":    adr_raw,
            })
    return persons


def inline_to_anchor(drawing_elem, pos_h_emu, pos_v_emu, width_emu, height_emu, dp_id):
    inline = drawing_elem.find(qn("wp:inline"))
    if inline is None:
        return

    docPr   = inline.find(qn("wp:docPr"))
    graphic = inline.find(qn("a:graphic"))
    cNvGFP  = inline.find(qn("wp:cNvGraphicFramePr"))

    anchor = OxmlElement("wp:anchor")
    for attr, val in [("behindDoc","1"), ("simplePos","0"), ("relativeHeight","0"),
                      ("locked","0"), ("layoutInCell","1"), ("allowOverlap","1"),
                      ("distT","0"), ("distB","0"), ("distL","0"), ("distR","0")]:
        anchor.set(attr, val)

    sp = OxmlElement("wp:simplePos"); sp.set("x","0"); sp.set("y","0")
    anchor.append(sp)

    ph = OxmlElement("wp:positionH"); ph.set("relativeFrom","page")
    oh = OxmlElement("wp:posOffset"); oh.text = str(pos_h_emu); ph.append(oh)
    anchor.append(ph)

    pv = OxmlElement("wp:positionV"); pv.set("relativeFrom","page")
    ov = OxmlElement("wp:posOffset"); ov.text = str(pos_v_emu); pv.append(ov)
    anchor.append(pv)

    ext = OxmlElement("wp:extent")
    ext.set("cx", str(width_emu)); ext.set("cy", str(height_emu))
    anchor.append(ext)

    anchor.append(OxmlElement("wp:wrapNone"))

    dp = deepcopy(docPr)
    dp.set("id", str(dp_id)); dp.set("name", f"Img{dp_id}")
    anchor.append(dp)

    if cNvGFP is not None:
        anchor.append(deepcopy(cNvGFP))

    g = deepcopy(graphic)
    pic_ns = "http://schemas.openxmlformats.org/drawingml/2006/picture"
    a_ns   = "http://schemas.openxmlformats.org/drawingml/2006/main"
    ext_el = g.find(f".//{{{a_ns}}}graphicData/{{{pic_ns}}}pic/"
                    f"{{{pic_ns}}}spPr/{{{a_ns}}}xfrm/{{{a_ns}}}ext")
    if ext_el is not None:
        ext_el.set("cx", str(width_emu))
        ext_el.set("cy", str(height_emu))
    anchor.append(g)

    drawing_elem.remove(inline)
    drawing_elem.append(anchor)


def make_vml_textbox_run(text, x_cm, y_cm, font_name, font_size_half_pt, shape_id):
    text_esc = (text.replace("&", "&amp;")
                    .replace("<", "&lt;")
                    .replace(">", "&gt;")
                    .replace('"', "&quot;"))

    style = (f"position:absolute;"
             f"margin-left:{x_cm:.2f}cm;margin-top:{y_cm:.2f}cm;"
             f"width:8cm;height:0.7cm;"
             f"z-index:1;"
             f"mso-position-horizontal-relative:page;"
             f"mso-position-vertical-relative:page")

    ns = f'xmlns:w="{W_NS}" xmlns:v="{V_NS}" xmlns:o="{O_NS}"'

    xml = f'''<w:r {ns}>
      <w:rPr><w:noProof/></w:rPr>
      <w:pict>
        <v:shape id="tb_{shape_id}" style="{style}"
                 type="#_x0000_t202" filled="f" stroked="f">
          <v:textbox inset="0,0,0,0">
            <w:txbxContent>
              <w:p>
                <w:pPr>
                  <w:spacing w:line="240" w:lineRule="auto"
                             w:before="0" w:after="0"/>
                  <w:rPr>
                    <w:rFonts w:ascii="{font_name}" w:hAnsi="{font_name}"
                              w:eastAsia="{font_name}" w:cs="{font_name}"/>
                    <w:sz w:val="{font_size_half_pt}"/>
                    <w:szCs w:val="{font_size_half_pt}"/>
                  </w:rPr>
                </w:pPr>
                <w:r>
                  <w:rPr>
                    <w:rFonts w:ascii="{font_name}" w:hAnsi="{font_name}"
                              w:eastAsia="{font_name}" w:cs="{font_name}"/>
                    <w:sz w:val="{font_size_half_pt}"/>
                    <w:szCs w:val="{font_size_half_pt}"/>
                    <w:color w:val="{TEXT_COLOR}"/>
                  </w:rPr>
                  <w:t xml:space="preserve">{text_esc}</w:t>
                </w:r>
              </w:p>
            </w:txbxContent>
          </v:textbox>
        </v:shape>
      </w:pict>
    </w:r>'''
    return etree.fromstring(xml)


def make_vml_textbox_run_rotated(text, target_x, target_y,
                                  font_name, font_size_half_pt, shape_id):
    text_esc = (text.replace("&", "&amp;")
                    .replace("<", "&lt;")
                    .replace(">", "&gt;")
                    .replace('"', "&quot;"))

    tb_w = 0.7
    tb_h = 10.0

    ml = target_x
    mt = target_y - tb_h

    style = (f"position:absolute;"
             f"margin-left:{ml:.2f}cm;margin-top:{mt:.2f}cm;"
             f"width:{tb_w:.1f}cm;height:{tb_h:.1f}cm;"
             f"z-index:1;"
             f"mso-position-horizontal-relative:page;"
             f"mso-position-vertical-relative:page")

    ns = f'xmlns:w="{W_NS}" xmlns:v="{V_NS}" xmlns:o="{O_NS}"'

    xml = f'''<w:r {ns}>
      <w:rPr><w:noProof/></w:rPr>
      <w:pict>
        <v:shape id="tb_{shape_id}" style="{style}"
                 type="#_x0000_t202" filled="f" stroked="f">
          <v:textbox style="layout-flow:vertical;mso-layout-flow-alt:bottom-to-top"
                     inset="0,0,0,0">
            <w:txbxContent>
              <w:p>
                <w:pPr>
                  <w:spacing w:line="240" w:lineRule="auto"
                             w:before="0" w:after="0"/>
                  <w:rPr>
                    <w:rFonts w:ascii="{font_name}" w:hAnsi="{font_name}"
                              w:eastAsia="{font_name}" w:cs="{font_name}"/>
                    <w:sz w:val="{font_size_half_pt}"/>
                    <w:szCs w:val="{font_size_half_pt}"/>
                  </w:rPr>
                </w:pPr>
                <w:r>
                  <w:rPr>
                    <w:rFonts w:ascii="{font_name}" w:hAnsi="{font_name}"
                              w:eastAsia="{font_name}" w:cs="{font_name}"/>
                    <w:sz w:val="{font_size_half_pt}"/>
                    <w:szCs w:val="{font_size_half_pt}"/>
                    <w:color w:val="{TEXT_COLOR}"/>
                  </w:rPr>
                  <w:t xml:space="preserve">{text_esc}</w:t>
                </w:r>
              </w:p>
            </w:txbxContent>
          </v:textbox>
        </v:shape>
      </w:pict>
    </w:r>'''
    return etree.fromstring(xml)


# ── Images et pagination ─────────────────────────────────────────────────────

def preparer_images():
    """Retourne l'image de la carte et sa version tournee de 90 degres."""
    bg = PILImage.open(BG_IMAGE)
    card_img = bg.crop((0, 0, IMG_W_PX, IMG_H_PX))
    card_rot_img = card_img.rotate(90, expand=True)
    return card_img, card_rot_img


def _png_bytes(img):
    buf = BytesIO()
    img.save(buf, "PNG")
    buf.seek(0)
    return buf


def _decouper_pages(persons):
    """Decoupe la liste des personnes en pages de CARDS_PER_PAGE cartes."""
    return [persons[i:i + CARDS_PER_PAGE]
            for i in range(0, len(persons), CARDS_PER_PAGE)]


# ── Generation Word (.docx) ──────────────────────────────────────────────────

def _drawing_template(doc, png_buffer):
    """Insere l'image une fois pour recuperer son XML de dessin, puis la retire."""
    para = doc.add_paragraph()
    run = para.add_run()
    run.add_picture(png_buffer)
    drawing = deepcopy(run._r.find(qn("w:drawing")))
    doc._body._body.remove(para._p)
    return drawing


def generer_docx(persons, card_img, card_rot_img):
    """Genere le document Word (une page A4 par groupe de 6 cartes)."""
    doc = Document()

    section = doc.sections[0]
    section.page_width    = Cm(PAGE_W_CM)
    section.page_height   = Cm(PAGE_H_CM)
    section.left_margin   = Cm(SAFE_MARGIN_CM)
    section.right_margin  = Cm(SAFE_MARGIN_CM)
    section.top_margin    = Cm(SAFE_MARGIN_CM)
    section.bottom_margin = Cm(SAFE_MARGIN_CM)

    style = doc.styles["Normal"]
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after  = Pt(0)

    drawing_tmpl_normal  = _drawing_template(doc, _png_bytes(card_img))
    drawing_tmpl_rotated = _drawing_template(doc, _png_bytes(card_rot_img))

    normal_w_emu  = cm_emu(CARD_W_CM)
    normal_h_emu  = cm_emu(CARD_H_CM)
    rotated_w_emu = cm_emu(CARD_H_CM)
    rotated_h_emu = cm_emu(CARD_W_CM)

    pages = _decouper_pages(persons)
    dp_id    = 100
    shape_id = 1

    for page_idx, page_persons in enumerate(pages):
        para = doc.add_paragraph()
        para.paragraph_format.space_before = Pt(0)
        para.paragraph_format.space_after  = Pt(0)

        for slot, person in enumerate(page_persons[:LEFT_COUNT]):
            card_left = LEFT_X
            card_top  = LEFT_TOPS[slot]

            run_el = OxmlElement("w:r")
            draw   = deepcopy(drawing_tmpl_normal)
            inline_to_anchor(draw, cm_emu(card_left), cm_emu(card_top),
                             normal_w_emu, normal_h_emu, dp_id)
            dp_id += 1
            run_el.append(draw)
            para._p.append(run_el)

            for fname, fx, fy in FIELDS:
                val = person.get(fname, "")
                if not val:
                    continue
                field_x = card_left + px_to_cm_x(fx)
                field_y = card_top + px_to_cm_y(fy)
                tb = make_vml_textbox_run(val, field_x, field_y,
                                          FONT_NAME, FONT_SIZE_HALF_PT, shape_id)
                shape_id += 1
                para._p.append(tb)

        for slot, person in enumerate(page_persons[LEFT_COUNT:]):
            card_left = RIGHT_X
            card_top  = RIGHT_TOPS[slot]

            run_el = OxmlElement("w:r")
            draw   = deepcopy(drawing_tmpl_rotated)
            inline_to_anchor(draw, cm_emu(card_left), cm_emu(card_top),
                             rotated_w_emu, rotated_h_emu, dp_id)
            dp_id += 1
            run_el.append(draw)
            para._p.append(run_el)

            for fname, fx, fy in FIELDS:
                val = person.get(fname, "")
                if not val:
                    continue
                rot_x = fy * CARD_H_CM / IMG_H_PX
                rot_y = (IMG_W_PX - fx) * CARD_W_CM / IMG_W_PX
                tb = make_vml_textbox_run_rotated(val, card_left + rot_x, card_top + rot_y,
                                                  FONT_NAME, FONT_SIZE_HALF_PT, shape_id)
                shape_id += 1
                para._p.append(tb)

        if page_idx < len(pages) - 1:
            bp = doc.add_paragraph()
            bp.paragraph_format.space_before = Pt(0)
            bp.paragraph_format.space_after  = Pt(0)
            br = OxmlElement("w:br")
            br.set(qn("w:type"), "page")
            bp.add_run()._r.append(br)

    result = BytesIO()
    doc.save(result)
    result.seek(0)
    return result


# ── Generation PDF ───────────────────────────────────────────────────────────

def _pdf_y(y_cm_from_top):
    """Convertit une coordonnee verticale (cm depuis le haut) en points PDF."""
    return (PAGE_H_CM - y_cm_from_top) * CM


def _dessiner_texte_pdf(pdf, texte, x_cm, y_cm, tourne):
    """Ecrit un champ. Coordonnees en cm depuis le haut-gauche de la page."""
    taille = FONT_SIZE_EFF_PT
    pdf.setFont(PDF_FONT_NAME, taille)
    pdf.setFillColor(HexColor("#" + TEXT_COLOR))
    if tourne:
        # Texte lu de bas en haut : le point d'ancrage est le bas de la zone
        pdf.saveState()
        pdf.translate(x_cm * CM + taille * PDF_TEXT_ASCENT, _pdf_y(y_cm))
        pdf.rotate(90)
        pdf.drawString(0, 0, texte)
        pdf.restoreState()
    else:
        pdf.drawString(x_cm * CM, _pdf_y(y_cm) - taille * PDF_TEXT_ASCENT, texte)


def _dessiner_page_pdf(pdf, page_persons, carte, carte_tournee):
    for slot, person in enumerate(page_persons[:LEFT_COUNT]):
        top = LEFT_TOPS[slot]
        pdf.drawImage(carte, LEFT_X * CM, _pdf_y(top + CARD_H_CM),
                      CARD_W_CM * CM, CARD_H_CM * CM)
        for fname, fx, fy in FIELDS:
            val = person.get(fname, "")
            if val:
                _dessiner_texte_pdf(pdf, val, LEFT_X + px_to_cm_x(fx),
                                    top + px_to_cm_y(fy), tourne=False)

    for slot, person in enumerate(page_persons[LEFT_COUNT:]):
        top = RIGHT_TOPS[slot]
        pdf.drawImage(carte_tournee, RIGHT_X * CM, _pdf_y(top + CARD_W_CM),
                      CARD_H_CM * CM, CARD_W_CM * CM)
        for fname, fx, fy in FIELDS:
            val = person.get(fname, "")
            if val:
                abs_x = RIGHT_X + fy * CARD_H_CM / IMG_H_PX
                abs_y = top + (IMG_W_PX - fx) * CARD_W_CM / IMG_W_PX
                _dessiner_texte_pdf(pdf, val, abs_x, abs_y, tourne=True)


def generer_pdf(persons, card_img, card_rot_img):
    """Genere le PDF (memes positions que le Word)."""
    buf = BytesIO()
    pdf = pdf_canvas.Canvas(buf, pagesize=(PAGE_W_CM * CM, PAGE_H_CM * CM))
    pdf.setTitle("Cartes d'immatriculation")

    carte = ImageReader(card_img)
    carte_tournee = ImageReader(card_rot_img)

    for page_persons in _decouper_pages(persons):
        _dessiner_page_pdf(pdf, page_persons, carte, carte_tournee)
        pdf.showPage()

    pdf.save()
    buf.seek(0)
    return buf


# ── Routes ───────────────────────────────────────────────────────────────────

GENERATEURS = {"docx": generer_docx, "pdf": generer_pdf}

MIMETYPES = {
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pdf": "application/pdf",
}


@app.route("/")
def index():
    return render_template("index.html")


@app.route("/api/generate", methods=["POST"])
def generate():
    if "file" not in request.files:
        return jsonify({"error": "Aucun fichier envoye"}), 400

    file = request.files["file"]
    if file.filename == "":
        return jsonify({"error": "Aucun fichier selectionne"}), 400

    if not file.filename.lower().endswith((".xlsx", ".xls")):
        return jsonify({"error": "Le fichier doit etre un Excel (.xlsx)"}), 400

    fmt = request.args.get("format", "docx")
    if fmt not in GENERATEURS:
        return jsonify({"error": "Format demande inconnu"}), 400

    try:
        persons = lire_excel(BytesIO(file.read()))
        if not persons:
            raise ValueError("Aucune personne trouvee dans le fichier Excel.")

        card_img, card_rot_img = preparer_images()
        output = GENERATEURS[fmt](persons, card_img, card_rot_img)

    except ValueError as e:
        return jsonify({"error": str(e)}), 400
    except Exception as e:
        app.logger.exception("Echec de la generation %s", fmt)
        return jsonify({"error": f"Erreur lors de la generation : {str(e)}"}), 500

    original_name = os.path.splitext(file.filename)[0]
    return send_file(
        output,
        mimetype=MIMETYPES[fmt],
        as_attachment=True,
        download_name=f"{original_name}_cartes_remplies.{fmt}",
    )


@app.route("/api/template")
def download_template():
    """Telecharger le template Excel vierge."""
    template_path = os.path.join(BASE_DIR, "static", "Template_Immatriculation_Niger.xlsx")
    if os.path.exists(template_path):
        return send_file(template_path, as_attachment=True,
                         download_name="Template_Immatriculation_Niger.xlsx")
    return jsonify({"error": "Template non disponible"}), 404


@app.route("/api/font")
def download_font():
    """Telecharger la police Kingthings Trypewriter 2."""
    if os.path.exists(FONT_FILE):
        return send_file(FONT_FILE, as_attachment=True,
                         download_name="KingthingsTrypewriter2.ttf")
    return jsonify({"error": "Police non disponible"}), 404


if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5000))
    app.run(host="0.0.0.0", port=port, debug=False)
